"""Moteur de calcul, import du grand livre et export Excel (mêmes formules que le fichier d'origine)."""
import csv, hashlib, hmac, json, os, re, unicodedata
from datetime import datetime, date

MOIS = ["JAN", "FEV", "MARS", "AVR", "MAI", "JUIN", "JUIL", "AOUT", "SEPT", "OCT", "NOV", "DEC"]


# ---------------------------------------------------------------- utilitaires
def parse_num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[\s\u00a0\u202f]", "", str(v))
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace("FCFA", "").replace("XOF", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -x if neg else x


def parse_date(v):
    """Retourne 'jj/mm/aaaa' ou '' si non reconnue."""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    s = str(v or "").strip()[:10]
    for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, f).strftime("%d/%m/%Y")
        except ValueError:
            pass
    return ""


def fmt_num(x, blank_zero=True):
    if blank_zero and abs(x) < 0.005:
        return ""
    return f"{x:,.0f}".replace(",", "\u202f")


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return s.lower().strip()


# ---------------------------------------------------------------- état
def new_state():
    t = date.today()
    return {
        "meta": {"societe": "", "banque": "", "compte": "", "numero": "",
                 "mois": f"{MOIS[t.month - 1]} {t.year}", "au": t.strftime("%d/%m/%Y"),
                 "solde_compta": "DEBITEUR"},
        "report": {"compta": {"montant": 0.0, "sens": "Débit"}, "banque": {"montant": 0.0, "sens": "Crédit"}},
        "gl": [],    # {date, lib, debit, credit, pointe}
        "bank": [],  # {date, lib, debit, credit}
        "lock": None,    # {au, mois, date} quand le rapprochement est validé
        "pwd": None,     # {salt, hash} du mot de passe de déverrouillage
        "journal": [],
    }


def _report_row(state, key, label):
    r = state["report"][key]
    m = parse_num(r["montant"])
    return ("", f"{label} {state['meta']['au']}", m if r["sens"] == "Débit" else 0.0, m if r["sens"] == "Crédit" else 0.0)


def left_rows(state):
    """Compte banque dans nos livres = report + écritures du grand livre NON pointées."""
    rows = [_report_row(state, "compta", "Report solde au")]
    rows += [(g["date"], g["lib"], g["debit"], g["credit"]) for g in state["gl"] if not g.get("pointe")]
    return rows


def right_rows(state):
    """Compte dans les livres de la banque = report + lignes du relevé saisies."""
    rows = [_report_row(state, "banque", "Report solde au")]
    rows += [(b["date"], b["lib"], b["debit"], b["credit"]) for b in state["bank"]]
    return rows


def calc(left, right):
    """Reproduit les formules de la feuille (lignes 51 à 57)."""
    D51, E51 = sum(r[2] for r in left), sum(r[3] for r in left)
    H51, I51 = sum(r[2] for r in right), sum(r[3] for r in right)
    D52 = E51 - D51 if D51 - E51 < 0 else 0      # =IF(D51-E51<0,E51-D51,0)
    E52 = D51 - E51 if D51 - E51 > 0 else 0      # =IF(D51-E51>0,D51-E51,0)
    H52 = I51 - H51 if H51 - I51 < 0 else 0      # =IF(H51-I51<0,I51-H51,0)
    I52 = H51 - I51 if H51 - I51 > 0 else 0      # =IF(H51-I51>0,H51-I51,0)
    D53, E53, H53, I53 = D51 + D52, E51 + E52, H51 + H52, I51 + I52
    ok = lambda x: "OK" if abs(x) < 0.005 else x
    return dict(D51=D51, E51=E51, H51=H51, I51=I51, D52=D52, E52=E52, H52=H52, I52=I52,
                D53=D53, E53=E53, H53=H53, I53=I53,
                E54=ok(E52 - H52), E55=ok(D52 - I52), E56=ok(D53 - E53), E57=ok(H53 - I53))


# ---------------------------------------------------------------- import grand livre
def _read_rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
        txt = f.read()
    try:
        dialect = csv.Sniffer().sniff(txt[:4000], delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    return list(csv.reader(txt.splitlines(), dialect))


def import_grand_livre(path):
    """Lit un grand livre (xlsx/csv). Détecte les en-têtes Date / Libellé / Débit / Crédit."""
    data = _read_rows(path)
    kw = {"date": ("date",), "debit": ("debit",), "credit": ("credit",),
          "lib": ("libelle", "operation", "designation", "description", "intitule", "piece", "reference", "tiers")}
    hdr, cols = None, {}
    for i, row in enumerate(data[:40]):
        found = {k: [j for j, c in enumerate(row) if any(w in _norm(c) for w in ws)] for k, ws in kw.items()}
        if found["date"] and found["debit"] and found["credit"]:
            hdr, cols = i, found
            break
    if hdr is None:  # repli : colonnes A..D = date, libellé, débit, crédit
        hdr, cols = -1, {"date": [0], "lib": [1], "debit": [2], "credit": [3]}
    out = []
    for row in data[hdr + 1:]:
        g = lambda j: row[j] if j < len(row) else None
        lib = " / ".join(str(g(j)).strip() for j in cols["lib"] if g(j) not in (None, "")) if cols["lib"] else ""
        d, c = parse_num(g(cols["debit"][0])), parse_num(g(cols["credit"][0]))
        dt = parse_date(g(cols["date"][0]))
        if (d == 0 and c == 0) or (not dt and (not lib or _norm(lib).startswith(("total", "solde", "report")))):
            continue
        out.append({"date": dt, "lib": lib, "debit": d, "credit": c, "pointe": False})
    return out


# ---------------------------------------------------------------- fichiers projet
def save_project(state, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=1)


def load_project(path):
    with open(path, "r", encoding="utf-8") as f:
        s = json.load(f)
    base = new_state()
    base["meta"].update(s.get("meta", {}))
    for k in ("compta", "banque"):
        base["report"][k].update(s.get("report", {}).get(k, {}))
    base["gl"], base["bank"] = s.get("gl", []), s.get("bank", [])
    base["lock"], base["pwd"], base["journal"] = s.get("lock"), s.get("pwd"), s.get("journal", [])
    return base


# ---------------------------------------------------------------- export Excel (avec formules)
def export_xlsx(state, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Etat de rapprochement"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHI", (0.9, 14.7, 44.6, 15.1, 19.3, 14.7, 39.1, 19, 18.3)):
        ws.column_dimensions[col].width = w
    F = lambda **k: Font(name="Arial", **k)
    fill = lambda c: PatternFill("solid", fgColor=c)
    thin, dbl = Side(style="thin"), Side(style="double")
    m = state["meta"]
    ws.merge_cells("B2:I2")
    ws["B2"] = f"RAPPROCHEMENT BANCAIRE AU {m['au']}"
    ws["B2"].font, ws["B2"].alignment = F(bold=True, size=12), Alignment(horizontal="center")
    for cell, txt in (("B4", f"SOCIETE: {m['societe']}"), ("D4", f"CODE: {m['banque']}"), ("F4", f"COMPTE: {m['compte']}"),
                      ("H4", "AU"), ("B5", f"MOIS: {m['mois']}"), ("D5", f"COMPTE N°: {m['numero']}"),
                      ("G5", f"SOLDE COMPTABILITE: {m['solde_compta']}")):
        ws[cell], ws[cell].font = txt, F(bold=True, size=10)
    ws["I4"], ws["I4"].number_format, ws["I4"].font = datetime.strptime(m["au"], "%d/%m/%Y") if parse_date(m["au"]) else m["au"], "dd/mm/yyyy", F(bold=True, size=10)
    ws.merge_cells("B7:E7"); ws.merge_cells("F7:I7")
    ws["B7"], ws["F7"] = "COMPTE BANQUE DANS NOS LIVRES", f"COMPTE {m['compte'].upper()} DANS LES LIVRES DE LA BANQUE"
    for c, col in (("B7", "FFC000"), ("F7", "FFFF00")):
        ws[c].fill, ws[c].font, ws[c].alignment = fill(col), F(bold=True, size=10), Alignment(horizontal="center")
    ws.merge_cells("B8:C8"); ws.merge_cells("F8:G8")
    for c, t in (("B8", "Opérations"), ("D8", "Débit"), ("E8", "Crédit"), ("F8", "Opérations"), ("H8", "Débit"), ("I8", "Crédit")):
        ws[c], ws[c].font, ws[c].alignment = t, F(bold=True, size=12), Alignment(horizontal="center")
    L, R = left_rows(state), right_rows(state)
    last = 8 + max(len(L), len(R))
    for rows, cols in ((L, "BCDE"), (R, "FGHI")):
        for i, (d, lib, deb, cre) in enumerate(rows):
            r = 9 + i
            ws[f"{cols[0]}{r}"] = datetime.strptime(d, "%d/%m/%Y") if parse_date(d) else d
            ws[f"{cols[0]}{r}"].number_format = "dd/mm/yyyy"
            ws[f"{cols[1]}{r}"] = lib
            for c, v in ((cols[2], deb), (cols[3], cre)):
                if v:
                    ws[f"{c}{r}"] = v
            for c in cols:
                ws[f"{c}{r}"].font = F(size=11, bold=(i == 0))
    for r in range(9, last + 1):
        for c in "DEHI":
            ws[f"{c}{r}"].number_format = "#,##0"
    s, so, t = last + 2, last + 3, last + 4  # SOUS TOTAUX / SOLDE RAPPROCHE / TOTAUX
    for r, lab in ((s, "SOUS TOTAUX"), (so, "SOLDE RAPPROCHE"), (t, "TOTAUX")):
        for c in "BF":
            ws[f"{c}{r}"], ws[f"{c}{r}"].font = lab, F(bold=True, size=11)
    for c in "DEHI":
        ws[f"{c}{s}"] = f"=SUM({c}9:{c}{last + 1})"
    ws[f"D{so}"] = f"=IF(D{s}-E{s}<0,E{s}-D{s},0)"
    ws[f"E{so}"] = f"=IF(D{s}-E{s}>0,D{s}-E{s},0)"
    ws[f"H{so}"] = f"=IF(H{s}-I{s}<0,I{s}-H{s},0)"
    ws[f"I{so}"] = f"=IF(H{s}-I{s}>0,H{s}-I{s},0)"
    for c in "DEHI":
        ws[f"{c}{t}"] = f"=SUM({c}{s}:{c}{so})"
    colors = {"D": ("5B9BD5", "5B9BD5"), "E": ("FF0000",) * 2, "H": ("FF0000",) * 2, "I": ("5B9BD5",) * 2}
    for c in "DEHI":
        ws[f"{c}{s}"].font = F(bold=True, size=10)
        ws[f"{c}{so}"].fill, ws[f"{c}{so}"].font = fill(colors[c][0]), F(bold=True, size=10, color="FFFFFF")
        ws[f"{c}{t}"].fill = fill("92D050" if c in "DE" else "7030A0")
        ws[f"{c}{t}"].font = F(bold=True, size=10, color="000000" if c in "DE" else "FFFFFF")
        for r in (s, so, t):
            ws[f"{c}{r}"].number_format = "#,##0"
            ws[f"{c}{r}"].border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws[f"E{t + 1}"] = f'=IF(E{so}-H{so}=0,"OK",E{so}-H{so})'
    ws[f"E{t + 2}"] = f'=IF(D{so}-I{so}=0,"OK",D{so}-I{so})'
    ws[f"E{t + 3}"] = f'=IF(D{t}-E{t}=0,"OK",D{t}-E{t})'
    ws[f"E{t + 4}"] = f'=IF(H{t}-I{t}=0,"OK",H{t}-I{t})'
    for k, col in enumerate(("FF0000", "5B9BD5", "92D050", "7030A0")):
        ws[f"E{t + 1 + k}"].font, ws[f"E{t + 1 + k}"].number_format = F(bold=True, size=11, color=col), "#,##0"
    ws.merge_cells(f"C{t + 2}:C{t + 3}")
    ws[f"C{t + 2}"], ws[f"C{t + 2}"].font, ws[f"C{t + 2}"].fill = "CONTROLES OBLIGATOIRES", F(bold=True, size=12), fill("BFBFBF")
    ws[f"C{t + 2}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)


# ---------------------------------------------------------------- validation / verrouillage
def log(state, msg):
    state.setdefault("journal", []).append(f"{datetime.now():%d/%m/%Y %H:%M} — {msg}")


def _hash(pwd, salt):
    return hashlib.pbkdf2_hmac("sha256", pwd.encode("utf-8"), bytes.fromhex(salt), 200_000).hex()


def set_password(state, pwd):
    salt = os.urandom(16).hex()
    state["pwd"] = {"salt": salt, "hash": _hash(pwd, salt)}


def check_password(state, pwd):
    p = state.get("pwd")
    return bool(p) and hmac.compare_digest(_hash(pwd, p["salt"]), p["hash"])


def validate_check(state):
    """Un rapprochement ne peut être validé que si la date est valide et les 4 contrôles sont à OK."""
    if not parse_date(state["meta"]["au"]):
        return False, "La date « Au » de l'état est invalide (format jj/mm/aaaa)."
    c = calc(left_rows(state), right_rows(state))
    bad = [(k, c[k]) for k in ("E54", "E55", "E56", "E57") if c[k] != "OK"]
    if bad:
        return False, ("Les contrôles obligatoires ne sont pas tous à « OK » : le rapprochement n'est pas équilibré.\n\n"
                       + "\n".join(f"{k} : {fmt_num(v, False)}" for k, v in bad))
    return True, ""


def lock_state(state):
    m = state["meta"]
    state["lock"] = {"au": m["au"], "mois": m["mois"], "date": datetime.now().strftime("%d/%m/%Y %H:%M")}
    log(state, f"Rapprochement validé au {m['au']} — pointages et saisies verrouillés")


def unlock_state(state):
    au = (state.get("lock") or {}).get("au", "")
    state["lock"] = None
    log(state, f"Déverrouillage (mot de passe) du rapprochement au {au}")


def next_period(state):
    """Nouveau dossier : mêmes références, suspens non pointés et lignes du relevé reportés."""
    n = new_state()
    for k in ("societe", "banque", "compte", "numero", "solde_compta"):
        n["meta"][k] = state["meta"][k]
    n["gl"] = [dict(g, pointe=False) for g in state["gl"] if not g.get("pointe")]
    n["bank"] = [dict(b) for b in state["bank"]]
    n["pwd"] = state.get("pwd")
    log(n, f"Nouvelle période créée depuis le rapprochement au {state['meta']['au']}")
    return n


# ---------------------------------------------------------------- modèle d'importation
def export_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    hdr = ["Date", "N° pièce", "Libellé", "Débit", "Crédit"]

    def head(ws):
        ws.append(hdr)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(name="Arial", bold=True), PatternFill("solid", fgColor="FFC000"), Alignment(horizontal="center")
        for col, w in zip("ABCDE", (13, 16, 55, 18, 18)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Grand livre"
    head(ws)
    for r in range(2, 1002):
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = "#,##0"
    ex = wb.create_sheet("Exemple")
    head(ex)
    for d, p, l, db, cr in ((datetime(2025, 12, 3), "CHQ0001", "REGLEMENT FACTURE FOURNISSEUR", None, 850000),
                            (datetime(2025, 12, 8), "REM0012", "REMISE CHEQUES CLIENTS", 2500000, None),
                            (datetime(2025, 12, 15), "VIR0005", "VIREMENT SALAIRES", None, 4100000),
                            (datetime(2025, 12, 29), "AGIOS", "AGIOS ET FRAIS DE TENUE DE COMPTE", None, 40000)):
        ex.append([d, p, l, db, cr])
    for r in range(2, 6):
        ex[f"A{r}"].number_format = "dd/mm/yyyy"
        ex[f"D{r}"].number_format = ex[f"E{r}"].number_format = "#,##0"
    ins = wb.create_sheet("Instructions")
    ins.column_dimensions["A"].width = 110
    for line in ("MODÈLE D'IMPORTATION DU GRAND LIVRE BANQUE", "",
                 "1. Éditez le grand livre du compte banque (521 SYSCOHADA / 512 PCG) dans votre logiciel comptable.",
                 "2. Copiez les écritures dans la feuille « Grand livre », sous la ligne d'en-têtes (seule cette 1re feuille est importée).",
                 "3. Colonnes : Date (jj/mm/aaaa), N° pièce (facultatif), Libellé, Débit, Crédit. Un seul montant par ligne.",
                 "4. Ne copiez pas les lignes de total, de solde ou de report (le report se saisit sur la page d'accueil).",
                 "5. Gardez les en-têtes sur la première ligne. Un fichier .csv avec les mêmes colonnes est aussi accepté.",
                 "6. Dans le logiciel : page « Import du grand livre banque » > « Importer le grand livre… », puis pointez chaque écriture.",
                 "", "La feuille « Exemple » montre le format attendu ; elle n'est pas importée."):
        ins.append([line])
    ins["A1"].font = Font(name="Arial", bold=True, size=13)
    wb.save(path)


# ---------------------------------------------------------------- impression
def build_print_html(state):
    from html import escape as e
    m = state["meta"]
    L, R = left_rows(state), right_rows(state)
    c = calc(L, R)

    def rows(lst):
        return "".join(f"<tr{' class=rep' if i == 0 else ''}><td class=c>{e(d)}</td><td>{e(lib)}</td>"
                       f"<td class=n>{fmt_num(db)}</td><td class=n>{fmt_num(cr)}</td></tr>" for i, (d, lib, db, cr) in enumerate(lst))

    def table(title, bg, lst, a, b, sol, tot, tot_fg):
        return (f"<div class=half><div class=th style='background:{bg}'>{e(title)}</div><table>"
                "<tr><th class=c>Date</th><th>Opérations</th><th class=n>Débit</th><th class=n>Crédit</th></tr>"
                f"{rows(lst)}"
                f"<tr class=sub><td colspan=2>SOUS TOTAUX</td><td class=n>{fmt_num(c[a + '51'], False)}</td><td class=n>{fmt_num(c[b + '51'], False)}</td></tr>"
                f"<tr class=sub><td colspan=2>SOLDE RAPPROCHE</td><td class=n style='background:{sol[0]};color:#fff'>{fmt_num(c[a + '52'], False)}</td>"
                f"<td class=n style='background:{sol[1]};color:#fff'>{fmt_num(c[b + '52'], False)}</td></tr>"
                f"<tr class=sub><td colspan=2>TOTAUX</td><td class=n style='background:{tot};color:{tot_fg}'>{fmt_num(c[a + '53'], False)}</td>"
                f"<td class=n style='background:{tot};color:{tot_fg}'>{fmt_num(c[b + '53'], False)}</td></tr></table></div>")

    ctl = "".join(f"<tr><td>{lab}</td><td class=n style='color:{col};font-weight:bold'>{'OK' if c[k] == 'OK' else fmt_num(c[k], False)}</td></tr>"
                  for k, col, lab in (("E54", "#FF0000", "Solde rapproché Crédit (nos livres) − Débit (banque)"),
                                      ("E55", "#5B9BD5", "Solde rapproché Débit (nos livres) − Crédit (banque)"),
                                      ("E56", "#5E9E2A", "Totaux nos livres : Débit − Crédit"),
                                      ("E57", "#7030A0", "Totaux banque : Débit − Crédit")))
    lk = state.get("lock")
    status = (f"<p class=ok>RAPPROCHEMENT VALIDÉ le {e(lk['date'])} (période au {e(lk['au'])}) — pointages verrouillés</p>" if lk
              else "<p class=draft>Rapprochement en cours — non validé</p>")
    left = table("COMPTE BANQUE DANS NOS LIVRES", "#FFC000", L, "D", "E", ("#5B9BD5", "#FF0000"), "#92D050", "#000")
    right = table(f"COMPTE {m['compte'].upper()} DANS LES LIVRES DE LA BANQUE", "#FFFF00", R, "H", "I", ("#FF0000", "#5B9BD5"), "#7030A0", "#fff")
    return f"""<!DOCTYPE html><html lang="fr"><head><meta charset="utf-8"><title>Rapprochement bancaire au {e(m['au'])}</title>
<style>
@page {{ size: A4 landscape; margin: 10mm; }}
* {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
body {{ font-family: Arial, sans-serif; font-size: 10pt; margin: 0; }}
h1 {{ text-align: center; font-size: 15pt; margin: 4px 0 8px; }}
td.l {{ border: 0; vertical-align: top; padding: 0 4px; }} table.info td {{ border: 0; font-weight: bold; padding: 2px 14px 6px 0; }}
.th {{ text-align: center; font-weight: bold; padding: 4px; border: 1px solid #000; }}
table {{ border-collapse: collapse; width: 100%; }}
td, th {{ border: 1px solid #000; padding: 2px 5px; font-size: 9pt; }} th {{ background: #eee; }}
.n {{ text-align: right; white-space: nowrap; }} .c {{ text-align: center; white-space: nowrap; }}
tr.rep td, tr.sub td {{ font-weight: bold; }}
.ctl th {{ background: #BFBFBF; font-size: 11pt; }} table.sign td {{ height: 80px; vertical-align: top; font-weight: bold; width: 33%; }}
.ok {{ color: #1e7b1e; font-weight: bold; text-align: center; }} .draft {{ color: #a00; font-weight: bold; text-align: center; }}
</style></head><body onload="window.print()">
<h1>RAPPROCHEMENT BANCAIRE AU …{e(m['au'])}…</h1>
<table class="info"><tr><td>SOCIETE : {e(m['societe'])}</td><td>CODE : {e(m['banque'])}</td><td>COMPTE : {e(m['compte'])}</td>
<td>COMPTE N° : {e(m['numero'])}</td><td>MOIS : {e(m['mois'])}</td><td>SOLDE COMPTABILITE : {e(m['solde_compta'])}</td></tr></table>
<table><tr><td class="l" width="50%">{left}</td><td class="l" width="50%">{right}</td></tr></table>
<table style="margin-top:10px"><tr><td class="l" width="46%"><table class="ctl"><tr><th colspan=2>CONTROLES OBLIGATOIRES</th></tr>{ctl}</table></td>
<td class="l"><table class="sign"><tr><td>Établi par :<br>Date :</td><td>Vérifié par :<br>Date :</td><td>Approuvé par :<br>Date :</td></tr></table></td></tr></table>
{status}
</body></html>"""


# ---------------------------------------------------------------- import cumulatif / modification
def _gl_key(g):
    return (g["date"], g["lib"].strip().lower(), round(g["debit"], 2), round(g["credit"], 2))


def merge_import(existing, new):
    """Ajoute `new` à la suite de `existing` sans rien écraser.
    Une écriture strictement identique (date, libellé, montants) déjà présente est ignorée,
    en tenant compte du nombre d'occurrences (deux chèques identiques le même jour restent possibles).
    Retourne (lignes_à_ajouter, nombre_ignoré)."""
    from collections import Counter
    have = Counter(_gl_key(g) for g in existing)
    added, skipped = [], 0
    for g in new:
        k = _gl_key(g)
        if have[k] > 0:
            have[k] -= 1
            skipped += 1
        else:
            added.append(g)
    return added, skipped


def edit_gl_line(state, i, date_s, lib, debit, credit):
    """Modifie une écriture du grand livre. Retourne un message d'erreur, ou None si OK."""
    d = parse_date(date_s)
    if not d:
        return "Date invalide (format jj/mm/aaaa)."
    if debit < 0 or credit < 0:
        return "Les montants doivent être positifs."
    if debit == 0 and credit == 0:
        return "Saisissez un montant au débit ou au crédit."
    if debit and credit:
        return "Un seul montant par ligne : débit OU crédit."
    g = state["gl"][i]
    old = f"{g['date']} | {g['lib']} | D {fmt_num(g['debit'], False)} | C {fmt_num(g['credit'], False)}"
    g.update(date=d, lib=lib.strip(), debit=debit, credit=credit, modifie=True)
    log(state, f"Écriture modifiée : {old}  →  {d} | {lib.strip()} | D {fmt_num(debit, False)} | C {fmt_num(credit, False)}")
    return None
