import os, sys, unittest
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import core


class TestCore(unittest.TestCase):
    def test_formules_equilibre(self):
        s = core.new_state()
        s["report"]["compta"] = {"montant": 1000, "sens": "Crédit"}
        s["report"]["banque"] = {"montant": 1000, "sens": "Crédit"}
        r = core.calc(core.left_rows(s), core.right_rows(s))
        self.assertEqual((r["D51"], r["E51"], r["E52"], r["D52"]), (0, 1000, 0, 1000))
        self.assertEqual((r["D53"], r["E53"]), (1000, 1000))
        self.assertEqual(r["E56"], "OK")
        self.assertEqual(r["E57"], "OK")

    def test_pointage(self):
        s = core.new_state()
        s["gl"] = [{"date": "01/01/2025", "lib": "A", "debit": 100, "credit": 0, "pointe": True},
                   {"date": "02/01/2025", "lib": "B", "debit": 0, "credit": 40, "pointe": False}]
        self.assertEqual(len(core.left_rows(s)), 2)  # report + 1 non pointée
        r = core.calc(core.left_rows(s), core.right_rows(s))
        self.assertEqual(r["E51"], 40)

    def test_nombres_dates(self):
        self.assertEqual(core.parse_num("1 234 567,50"), 1234567.5)
        self.assertEqual(core.parse_num("1,234,567"), 1234567)
        self.assertEqual(core.parse_num("(500)"), -500)
        self.assertEqual(core.parse_date("2025-12-31"), "31/12/2025")

    def _balanced(self):
        s = core.new_state()
        s["meta"]["au"] = "31/12/2025"
        s["report"]["compta"] = {"montant": 1000, "sens": "Débit"}
        s["report"]["banque"] = {"montant": 1000, "sens": "Crédit"}
        return s

    def test_validation_et_mot_de_passe(self):
        s = self._balanced()
        self.assertTrue(core.validate_check(s)[0])
        s["bank"].append({"date": "01/12/2025", "lib": "frais", "debit": 0, "credit": 50})
        self.assertFalse(core.validate_check(s)[0])          # déséquilibré -> refus
        s["bank"] = []
        core.set_password(s, "secret"); core.lock_state(s)
        self.assertTrue(s["lock"]); self.assertTrue(core.check_password(s, "secret")); self.assertFalse(core.check_password(s, "x"))
        core.unlock_state(s)
        self.assertIsNone(s["lock"]); self.assertEqual(len(s["journal"]), 2)

    def test_sauvegarde_verrou_et_nouvelle_periode(self):
        import tempfile
        s = self._balanced()
        s["gl"] = [{"date": "01/12/2025", "lib": "A", "debit": 5, "credit": 0, "pointe": True},
                   {"date": "02/12/2025", "lib": "B", "debit": 0, "credit": 7, "pointe": False}]
        core.set_password(s, "pw"); core.lock_state(s)
        p = os.path.join(tempfile.mkdtemp(), "x.rapp"); core.save_project(s, p)
        t = core.load_project(p)["accounts"][0]      # un état seul est enveloppé dans un projet
        self.assertTrue(t["lock"]); self.assertTrue(core.check_password(t, "pw"))
        n = core.next_period(t)
        self.assertEqual((n["meta"]["au"], n["meta"]["mois"]), ("31/01/2026", "JAN 2026"))
        s2 = core.new_state(); s2["meta"]["au"] = "30/12/2026"; self.assertEqual(core.next_period(s2)["meta"]["au"], "31/01/2027")
        s2["meta"]["au"] = "31/01/2026"; self.assertEqual(core.next_period(s2)["meta"]["au"], "28/02/2026")
        self.assertEqual([g["lib"] for g in n["gl"]], ["B"]); self.assertIsNone(n["lock"]); self.assertTrue(core.check_password(n, "pw"))

    def test_modele_et_impression(self):
        import tempfile
        p = os.path.join(tempfile.mkdtemp(), "m.xlsx"); core.export_template(p)
        self.assertEqual(core.import_grand_livre(p), [])      # modèle vide : rien à importer
        from openpyxl import load_workbook
        wb = load_workbook(p); ws = wb["Grand livre"]
        ws.append(["03/12/2025", "CHQ1", "REGLEMENT <FOURNISSEUR>", None, 850000]); wb.save(p)
        r = core.import_grand_livre(p)
        self.assertEqual((r[0]["date"], r[0]["credit"]), ("03/12/2025", 850000))
        s = core.new_state(); s["gl"] = r
        h = core.build_print_html(s)
        self.assertIn("850", h); self.assertIn("&lt;FOURNISSEUR&gt;", h); self.assertIn("non validé", h)

    def test_import_cumulatif_et_modification(self):
        s = core.new_state()
        jan = [{"date": "05/01/2026", "lib": "CHQ A", "debit": 0, "credit": 100, "pointe": False},
               {"date": "06/01/2026", "lib": "CHQ B", "debit": 0, "credit": 50, "pointe": False},
               {"date": "06/01/2026", "lib": "CHQ B", "debit": 0, "credit": 50, "pointe": False}]   # 2 lignes identiques légitimes
        add, sk = core.merge_import(s["gl"], jan); s["gl"] += add
        self.assertEqual((len(add), sk), (3, 0))
        s["gl"][0]["pointe"] = True
        fev = [{"date": "03/02/2026", "lib": "REMISE C", "debit": 70, "credit": 0, "pointe": False}]
        add, sk = core.merge_import(s["gl"], fev); s["gl"] += add
        self.assertEqual(len(s["gl"]), 4); self.assertTrue(s["gl"][0]["pointe"])          # rien écrasé, pointage conservé
        add, sk = core.merge_import(s["gl"], jan + fev)                                   # ré-import : tout est déjà là
        self.assertEqual((len(add), sk), (0, 4))
        self.assertIsNone(core.edit_gl_line(s, 1, "07/01/2026", "CHQ B corrigé", 0, 55))
        self.assertEqual((s["gl"][1]["credit"], s["gl"][1]["date"]), (55, "07/01/2026")); self.assertIn("Écriture modifiée", s["journal"][-1])
        self.assertIn("invalide", core.edit_gl_line(s, 1, "xx", "a", 0, 5))
        self.assertIn("débit OU crédit", core.edit_gl_line(s, 1, "07/01/2026", "a", 5, 5))

    def test_plusieurs_banques(self):
        import tempfile
        from openpyxl import load_workbook
        pr = core.new_project()
        a = pr["accounts"][0]; a["meta"].update(societe="ACME", banque="BOA", compte="Compte courant", au="31/12/2025", mois="DEC 2025")
        a["report"] = {"compta": {"montant": 1000, "sens": "Débit"}, "banque": {"montant": 1000, "sens": "Crédit"}}
        a["gl"] = [{"date": "01/12/2025", "lib": "A", "debit": 5, "credit": 0, "pointe": True},
                   {"date": "02/12/2025", "lib": "B", "debit": 0, "credit": 7, "pointe": False}]
        core.set_password(a, "pw")
        b = core.add_account(pr, "CORIS")                                   # 2e banque
        self.assertEqual((len(pr["accounts"]), pr["cur"]), (2, 1))
        self.assertEqual((b["meta"]["societe"], b["meta"]["au"], b["meta"]["banque"]), ("ACME", "31/12/2025", "CORIS"))
        self.assertEqual(b["gl"], []); self.assertTrue(core.check_password(b, "pw"))
        b["report"] = {"compta": {"montant": 500, "sens": "Débit"}, "banque": {"montant": 500, "sens": "Crédit"}}
        self.assertEqual(core.account_label(a, 0), "BOA — Compte courant"); self.assertEqual(core.account_label(core.new_state(), 2), "Banque 3")
        core.set_societe(pr, "ACME SA"); self.assertEqual((a["meta"]["societe"], b["meta"]["societe"]), ("ACME SA", "ACME SA"))
        # indépendance des comptes
        self.assertNotEqual(core.calc(core.left_rows(a), core.right_rows(a))["D51"], core.calc(core.left_rows(b), core.right_rows(b))["D51"])
        self.assertEqual(core.unvalidated(pr), ["BOA — Compte courant", "CORIS"])
        a["bank"] = [{"date": "03/12/2025", "lib": "frais", "debit": 0, "credit": 7}]
        self.assertEqual(len(core.summary(pr)), 2)
        core.lock_state(b); core.set_societe(pr, "X"); self.assertEqual(b["meta"]["societe"], "ACME SA")   # compte validé figé
        # sauvegarde / réouverture
        d = tempfile.mkdtemp(); p = os.path.join(d, "m.rapp"); core.save_project(pr, p)
        q = core.load_project(p)
        self.assertEqual((len(q["accounts"]), q["cur"]), (2, 1)); self.assertEqual(q["accounts"][1]["meta"]["banque"], "CORIS")
        # période suivante : tous les comptes
        core.lock_state(a); self.assertEqual(core.unvalidated(pr), [])
        n = core.next_project(pr)
        self.assertEqual([x["meta"]["au"] for x in n["accounts"]], ["31/01/2026", "31/01/2026"])
        self.assertEqual([g["lib"] for g in n["accounts"][0]["gl"]], ["B"]); self.assertEqual(len(n["accounts"][0]["bank"]), 1)
        self.assertTrue(all(x["lock"] is None for x in n["accounts"]))
        # export Excel : une feuille par banque ; impression : une page par banque
        x = os.path.join(d, "e.xlsx"); core.export_xlsx(pr, x)
        self.assertEqual(load_workbook(x).sheetnames, ["BOA — Compte courant", "CORIS"])
        h = core.build_print_html(pr); self.assertEqual(h.count('class="pg"'), 2); self.assertIn("CORIS", h)
        self.assertEqual(core.build_print_html(a).count('class="pg"'), 1)
        # suppression d'une banque (la dernière ne peut pas être supprimée)
        self.assertTrue(core.remove_account(pr, 0)); self.assertFalse(core.remove_account(pr, 0)); self.assertEqual(pr["cur"], 0)

    def test_ancien_fichier_un_seul_compte(self):
        import tempfile, json
        p = os.path.join(tempfile.mkdtemp(), "old.rapp")
        old = core.new_state(); old["meta"]["banque"] = "BOA"
        with open(p, "w", encoding="utf-8") as f: json.dump(old, f)      # ancien format, sans « accounts »
        pr = core.load_project(p)
        self.assertEqual((len(pr["accounts"]), pr["cur"], pr["accounts"][0]["meta"]["banque"]), (1, 0, "BOA"))


if __name__ == "__main__":
    unittest.main()
